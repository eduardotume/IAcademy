"""
exportar.py — Convierte las respuestas de la IA (en Markdown) a un archivo Word (.docx).
"""

import io
import re
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor

MORADO = RGBColor(0x7C, 0x3A, 0xED)


def _agregar_con_negritas(parrafo, texto):
    """Respeta **negritas** y *cursivas* simples del Markdown."""
    partes = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto)
    for parte in partes:
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            parrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            parrafo.add_run(parte[1:-1]).italic = True
        elif parte:
            parrafo.add_run(parte)


def markdown_a_docx(titulo: str, contenido: str, autor: str = "") -> bytes:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    encabezado = doc.add_heading(titulo, level=0)
    for run in encabezado.runs:
        run.font.color.rgb = MORADO

    meta = f"Generado con IAcademy · {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    if autor:
        meta = f"{autor} · " + meta
    p = doc.add_paragraph(meta)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    for linea in contenido.splitlines():
        limpia = linea.strip()
        if not limpia:
            continue
        if limpia in ("---", "***", "___"):
            continue
        titulo_md = re.match(r"^(#{1,6})\s+(.*)", limpia)
        vineta = re.match(r"^[-*•]\s+(.*)", limpia)
        numerada = re.match(r"^\d+[.)]\s+(.*)", limpia)

        if titulo_md:
            nivel = min(len(titulo_md.group(1)), 3)
            doc.add_heading(titulo_md.group(2).replace("**", ""), level=nivel)
        elif vineta:
            _agregar_con_negritas(doc.add_paragraph(style="List Bullet"), vineta.group(1))
        elif numerada:
            _agregar_con_negritas(doc.add_paragraph(style="List Number"), numerada.group(1))
        else:
            _agregar_con_negritas(doc.add_paragraph(), limpia)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
